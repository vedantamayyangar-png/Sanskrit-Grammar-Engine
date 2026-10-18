#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_sarga1_report_xlsx.py
================================
Same data as generate_sarga1_report.py (the combined sandhi_splitter.py +
samasa_splitter.py pipeline run over every word of Balakanda Sarga 1),
written as an .xlsx workbook instead of a fixed-width text table -- one
row per word, real (unwrapped) cell text, plus two blank reviewer
columns ("Verdict" and "Correction/Notes") for marking errors by hand.

USAGE
-----
    D:\\python\\python.exe generate_sarga1_report_xlsx.py [output.xlsx]
    (default: sarga1_sandhi_samasa_report.xlsx, next to this script)
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SANDHI_DIR = os.path.normpath(os.path.join(HERE, "..", "sandhi"))
for d in (SANDHI_DIR, HERE):
    if d not in sys.path:
        sys.path.insert(0, d)

import openpyxl  # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from generate_sarga1_report import (  # noqa: E402
    build_verse_words, _sandhi_name, _sandhi_vigraha,
    _samasa_name, _samasa_vigraha,
)
from combined_splitter import CombinedSplitter  # noqa: E402

DEFAULT_OUTPUT = os.path.join(HERE, "sarga1_sandhi_samasa_report.xlsx")

HEADERS = ["#", "Word", "Verse", "Sandhi Rule(s)", "Sandhi Pada-Vigraha",
           "Samasa Type", "Samasa Pada-Vigraha", "Verdict (mark here)",
           "Correction / Notes"]
COL_WIDTHS = [5, 22, 14, 18, 28, 30, 28, 16, 40]


def main():
    output_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUTPUT

    cs = CombinedSplitter()
    verse_words = build_verse_words()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sarga 1 report"

    ws.append(HEADERS)
    for col, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    n_sandhi = 0
    n_samasa = 0
    for i, (ref, word) in enumerate(verse_words, 1):
        result = cs.split_word(word)
        sandhi_name = _sandhi_name(result["sandhi_rules"])
        sandhi_vigraha = _sandhi_vigraha(result["sandhi_pieces"])
        samasa_name = _samasa_name(result["samasa_decompositions"])
        samasa_vigraha = _samasa_vigraha(result["samasa_decompositions"])
        if sandhi_name != "-":
            n_sandhi += 1
        if samasa_name != "-":
            n_samasa += 1
        ws.append([i, word, ref, sandhi_name, sandhi_vigraha,
                   samasa_name, samasa_vigraha, "", ""])
        print(f"[{i}/{len(verse_words)}] {ref} {word}", file=sys.stderr)

    # Wrap long text cells (columns D-I) and align the row-number/verse
    # columns for readability once opened in Excel.
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column > 3))

    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Valmiki Ramayana -- Balakanda, Sarga 1 (R_1,001.001-079)"
    ws2["A2"] = "Combined sandhi_splitter.py + samasa_splitter.py analysis"
    ws2["A3"] = f"Total words: {len(verse_words)}  |  Words with a sandhi split: {n_sandhi}  |  Words with a samasa decomposition: {n_samasa}"
    ws2["A5"] = ("Basis: 'Sandhi Rule(s)'/'Sandhi Pada-Vigraha' come directly from "
                 "sandhi_splitter.py's own niyama-rule citations and dictionary "
                 "validation (apte_dictionary.sqlite). 'Samasa Type' is an explicitly "
                 "heuristic structural guess (samasa_splitter.py's guess_samasa_type, "
                 "NOT a substitute for real meaning-based classification); 'Samasa "
                 "Pada-Vigraha' is only populated when a dictionary-verified 2+-member "
                 "decomposition survived this pipeline's false-positive guards. A '-' "
                 "means 'checked, nothing found' -- never a guess presented as a result.")
    ws2["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 110
    ws2["A7"] = "Reviewer columns: mark 'Verdict' (e.g. OK / WRONG / PARTIAL) and add notes in 'Correction / Notes'."

    wb.save(output_path)
    print(f"\nWrote {output_path} ({len(verse_words)} words, "
          f"{n_sandhi} with sandhi splits, {n_samasa} with samasa decompositions)",
          file=sys.stderr)


if __name__ == "__main__":
    main()
