"""Assemble the analysis into: standalone openpyxl build script, the xlsx, and markdown tables."""
import importlib, glob, re, pprint, subprocess, os, sys
from translit import load_verses

OUT = "out"
os.makedirs(OUT, exist_ok=True)

vs = load_verses("sarga1_iast.txt")
ALL = {}
for f in sorted(glob.glob("d0*.py")):
    ALL.update(importlib.import_module(f[:-3]).V)


def samhita(k):
    """Join GRETIL's split tokens into continuous Devanagari (consonant-final + next, avagraha)."""
    lines = []
    for line in vs[k]["dev"].split("\n"):
        line = re.sub(r"\s*॥\s*\d+\s*॥", "", line).replace(" ।", "").strip()
        line = re.sub(r"्\s+", "्", line)          # वाल्मीकिर् मुनि → वाल्मीकिर्मुनि
        line = re.sub(r"\s+ऽ", "ऽ", line)          # ततो ऽगर्जत् → ततोऽगर्जत्
        lines.append(line)
    if len(lines) == 2:
        return f"{lines[0]} ।\n{lines[1]} ॥ {k} ॥"
    return f"{lines[0]} ।\n{lines[1]} ।\n{lines[2]} ॥ {k} ॥"


DATA = []
for k in range(1, 80):
    tr, rows = ALL[k]
    DATA.append({"num": k, "ref": f"R 1,001.{k:03d}", "verse": samhita(k),
                 "translation": tr, "rows": [list(r) for r in rows]})

HEADERS = [
    "Original Word (मूलपदम्)",
    "Sandhi Split & Paninian Sutra (सन्धिविच्छेदः सूत्रं च)",
    "Upasarga (उपसर्गः)",
    "Base Lemma / Pratipadika (प्रातिपदिकम्)",
    "Stripped Elements (विभक्तिः/विसर्गः/अव्ययम्)",
    "Samasa Analysis (समासविग्रहः प्रकारश्च)",
]

# ---------------------------------------------------------------- build script
SCRIPT = '''# -*- coding: utf-8 -*-
"""
Vālmīki Rāmāyaṇa — Bāla Kāṇḍa, Sarga 1 (79 ślokas)
Word-by-word Pāṇinian analysis → sanskrit_analysis.xlsx

Source text: GRETIL sa_rAmAyaNa.txt (critical edition, R_1,001.001–079), transliterated to Devanagari.
Run:  python build_sanskrit_analysis.py      (requires: pip install openpyxl)
"""
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sanskrit Grammatical Analysis"
ws.views.sheetView[0].showGridLines = True

# Style setup
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
cell_font = Font(name="Segoe UI", size=10)
word_font = Font(name="Segoe UI", size=11, bold=True)
section_font = Font(name="Segoe UI", size=12, bold=True)
section_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
trans_font = Font(name="Segoe UI", size=10, italic=True, color="404040")
trans_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center")
wrap_top = Alignment(wrap_text=True, vertical="top")

# Headers
headers = HEADERS_LITERAL

COL_WIDTHS = [22, 60, 16, 34, 58, 58]

DATA = DATA_LITERAL


def est_lines(text, width):
    """Rough wrapped-line estimate so Devanagari rows are tall enough (openpyxl cannot auto-fit)."""
    chars_per_line = max(8, int(width * 1.5))
    return sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\\n"))


for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.fill, c.font, c.border = header_fill, header_font, border
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[1].height = 36
ws.freeze_panes = "A2"

r = 2
for verse in DATA:
    # Section header row: verse number + Devanagari text
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value=f"श्लोकः {verse['num']}  ({verse['ref']})\\n{verse['verse']}")
    c.font, c.fill, c.alignment = section_font, section_fill, wrap
    for col in range(1, 7):
        ws.cell(row=r, column=col).border = border
    ws.row_dimensions[r].height = 18 * (verse["verse"].count("\\n") + 2) + 6
    r += 1

    # Translation row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="English Translation: " + verse["translation"])
    c.font, c.fill, c.alignment = trans_font, trans_fill, wrap
    for col in range(1, 7):
        ws.cell(row=r, column=col).border = border
    ws.row_dimensions[r].height = 15 * est_lines(verse["translation"], sum(COL_WIDTHS)) + 8
    r += 1

    # Word rows
    for row in verse["rows"]:
        lines = 1
        for col, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = word_font if col == 1 else cell_font
            c.alignment = wrap if col == 1 else wrap_top
            c.border = border
            lines = max(lines, est_lines(val, COL_WIDTHS[col - 1]))
        ws.row_dimensions[r].height = 16 * lines + 6
        r += 1

    r += 1  # blank spacer row between verses

# Adjust column dimensions to ensure no text clipping for Devanagari characters.
for i, w in enumerate(COL_WIDTHS, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws.auto_filter.ref = f"A1:F{r}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.print_title_rows = "1:1"

wb.save("sanskrit_analysis.xlsx")
print(f"Saved sanskrit_analysis.xlsx — {len(DATA)} verses, {sum(len(v['rows']) for v in DATA)} word rows")
'''

script = SCRIPT.replace("HEADERS_LITERAL", pprint.pformat(HEADERS, width=110)) \
               .replace("DATA_LITERAL", pprint.pformat(DATA, width=110, sort_dicts=False))
with open(f"{OUT}/build_sanskrit_analysis.py", "w", encoding="utf-8") as fh:
    fh.write(script)

subprocess.run([sys.executable, "build_sanskrit_analysis.py"], cwd=OUT, check=True)

# ---------------------------------------------------------------- markdown
md = ["# वाल्मीकिरामायणम् — बालकाण्डः, प्रथमः सर्गः",
      "",
      "Word-by-word Pāṇinian analysis of all 79 ślokas (GRETIL critical edition, R 1,001.001–079). "
      "Sandhi column shows the word as it meets its neighbour in the verse; where GRETIL's reading "
      "looks like a typo or an ārṣa form, the row says so.",
      ""]
for v in DATA:
    md += [f"## श्लोकः {v['num']}", "",
           "**Original Verse (Devanagari):**", "",
           "> " + v["verse"].replace("\n", "  \n> "), "",
           f"**English Translation:** {v['translation']}", "",
           "| " + " | ".join(HEADERS) + " |",
           "|" + "---|" * len(HEADERS)]
    for row in v["rows"]:
        md.append("| " + " | ".join(str(x).replace("|", "/") for x in row) + " |")
    md.append("")
with open(f"{OUT}/sanskrit_analysis_tables.md", "w", encoding="utf-8") as fh:
    fh.write("\n".join(md))
print("markdown written")
