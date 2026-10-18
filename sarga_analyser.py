#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sarga_analyser.py — word-by-word grammatical analysis of any GRETIL-format sarga, from the terminal,
using only the project's own tools and data (no internet, no LLM).

For every verse it writes the same six columns as the hand-written Sarga-1 treatise:

    मूलपदम् | सन्धिविच्छेदः सूत्रं च | उपसर्गः | प्रातिपदिकम् | विभक्तिः/विसर्गः/अव्ययम् | समासविग्रहः प्रकारश्च

plus a seventh column, "confidence", that says how each row was obtained (validated / rule /
guess), so a reader knows which cells to check by hand.

Pipeline (each stage is a separate, testable module):
  1. samasa_splitter.parse_gretil_verses             read the kāṇḍa/sarga from the GRETIL text
  2. combined_splitter.analyze_pada_line             junction sandhi between printed words
     (+ deterministic_junction fallback)             and words fused inside one printed token
     (+ split_fused_token)
  3. pada_analyser                                   prātipadika + vibhakti/vacana/liṅga, or
                                                     dhātu + lakāra/puruṣa/vacana/pada, or avyaya
  4. samasa_classifier                               samāsa members + type + reason
  5. citation_lookup                                 Devanagari sandhi/samāsa labels + sūtras

All modules, data files and outputs live in the same folder as this script.

Examples
  python sarga_analyser.py --kanda 1 --sarga 1
  python sarga_analyser.py --kanda 1 --sarga 2 --verses 1-10
  python sarga_analyser.py --kanda 1 --sarga 1 --gold .
  python sarga_analyser.py --source my_text.txt --kanda 2 --sarga 5 --formats xlsx,md,json
  python sarga_analyser.py --kanda 1 --sarga 1 --formats json --glossary meanings.json   # Padakosha JSON with meanings
  python sarga_analyser.py --kanda 1 --sarga 1 --formats json --json-style rows          # old per-verse JSON

Translations are not generated (there is no local translation engine); pass --translations
FILE.json ({"1": "…", "2": "…"}) to include your own.
"""
import argparse
import json
import os
import re
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from indic_transliteration import sanscript  # noqa: E402
from indic_transliteration.sanscript import transliterate as tl  # noqa: E402
from samasa_splitter import parse_gretil_verses  # noqa: E402
from combined_splitter import CombinedSplitter  # noqa: E402
from pada_analyser import NUM_NAMES  # noqa: E402
from padakosha_builder import PadakoshaBuilder  # noqa: E402

try:
    import citation_lookup as cite
except Exception:  # citation_lookup is optional: fall back to the raw IAST rule names
    cite = None

SUP = {(1, "e"): "सुँ", (1, "d"): "औ", (1, "b"): "जस्", (2, "e"): "अम्", (2, "d"): "औट्", (2, "b"): "शस्",
       (3, "e"): "टा", (3, "d"): "भ्याम्", (3, "b"): "भिस्", (4, "e"): "ङे", (4, "d"): "भ्याम्",
       (4, "b"): "भ्यस्", (5, "e"): "ङसिँ", (5, "d"): "भ्याम्", (5, "b"): "भ्यस्", (6, "e"): "ङस्",
       (6, "d"): "ओस्", (6, "b"): "आम्", (7, "e"): "ङि", (7, "d"): "ओस्", (7, "b"): "सुप्",
       (8, "e"): "सुँ", (8, "d"): "औ", (8, "b"): "जस्"}
HEADERS = ["Original Word (मूलपदम्)", "Sandhi Split & Paninian Sutra (सन्धिविच्छेदः सूत्रं च)",
           "Upasarga (उपसर्गः)", "Base Lemma / Pratipadika (प्रातिपदिकम्)",
           "Stripped Elements (विभक्तिः/विसर्गः/अव्ययम्)", "Samasa Analysis (समासविग्रहः प्रकारश्च)",
           "Confidence (प्रामाण्यम्)"]


def deva(iast):
    return tl(iast, sanscript.IAST, sanscript.DEVANAGARI)


def sandhi_label(rule, sutra):
    if cite is not None:
        label, citation, known = cite.lookup_sandhi_citation(rule)
        if known and label != "-":
            return f"{label} [{citation}]"
    return f"{rule} [{sutra}]"


def morph_cells(r):
    """(upasarga, prātipadika, stripped-elements, confidence) for one pada_analyser reading."""
    if r is None:
        return "नास्ति", "—", "अनिर्णीतम् (no reading found)", "none"
    if r["type"] == "noun":
        stem = r["stem"]
        ups = "नास्ति"
        if r.get("krt"):
            k = r["krt"]
            stem += f" (√{k['dhatu']} + {k['pratyaya']})"
            if k["upasarga"]:
                ups = ", ".join(k["upasarga"])
        stripped = (f"सुप्-प्रत्ययः «{SUP[(r['case'], r['num'])]}» — {r['case_name']} विभक्तिः, "
                    f"{NUM_NAMES[r['num']]}, {r['gender']}; {r['model']}-शब्दवत् (शब्दमञ्जरी पृ.{r['page']})")
        return ups, stem, stripped, r["confidence"]
    if r["type"] == "verb":
        stem = f"{r['dhatu']} (गणः {r['gana']}, {r['artha']})"
        stripped = f"तिङ् — {r['lakara']}-लकारः, {r['purusha']}, {r['vacana']}, {r['pada']}"
        if r.get("nijanta"):
            stripped += "; णिजन्तम्"
        if "आर्ष" in r["evidence"] or "यक्" in r["evidence"]:
            stripped += "; " + r["evidence"].split("; ", 1)[-1]
        conf = "validated" if r["_score"] >= 3.4 else "rule"
        return (", ".join(r["upasarga"]) or "नास्ति"), stem, stripped, conf
    # avyaya
    if r.get("kind") == "krdanta":
        stripped = f"कृदन्तम् अव्ययम् — कृत्-प्रत्ययः «{r['pratyaya']}» ({r.get('sutra', '')})"
        if r.get("nan"):
            stripped += "; नञ्-पूर्वकम्"
        return (", ".join(r.get("upasarga") or []) or "नास्ति"), r["stem"], stripped, "rule"
    if r.get("kind") == "taddhita":
        return "नास्ति", r["stem"], f"तद्धितान्तम् अव्ययम् — «{r['pratyaya']}»", "rule"
    return "नास्ति", r["stem"], "अव्ययम् — सुप्-लुक् (अव्ययादाप्सुपः)" + (f"; {r['gloss']}" if r.get("gloss") else ""), "validated"


def analyse_verse(cs, pa, lines):
    """Rows for one verse: list of dicts with the 7 output cells + raw readings."""
    out, raw = [], []
    for li, line in enumerate(lines):
        toks = merge_periphrastic(line.split())
        if not toks:
            continue
        rows = cs.analyze_pada_line(toks)
        raw.append({"tokens": toks, "rows": rows})
        for ti, row in enumerate(rows):
            last_in_line = ti == len(rows) - 1
            nxt = toks[ti + 1] if not last_in_line else None
            pieces = row["token_split"] or [row["citation_form"]]
            seps = [tuple(row["token_split_sandhi"])] if row["token_split"] else []
            for _ in range(2):          # a printed token may fuse three words: āpṛṣṭvā+eva+abhy…
                more = cs.split_fused_token(pieces[-1]) if len(pieces[-1]) > 5 else None
                if not more:
                    break
                pieces = pieces[:-1] + more["pieces"]
                seps.append(tuple(more["sandhi"]))
            for pi, piece in enumerate(pieces):
                is_last_piece = pi == len(pieces) - 1
                # --- sandhi cell
                if not is_last_piece:
                    styp, sutra = seps[pi]
                    s_cell = f"{deva(piece)} + {deva(pieces[pi + 1])} → {deva(row['word'])}: {sandhi_label(styp, sutra)}"
                elif row["sandhi_rule"]:
                    left, right = row["sandhi_vigraha"]
                    s_cell = (f"{deva(pieces[-1]) if row['token_split'] else deva(left)} + {deva(right)} → "
                              f"{deva(row['word'])} {deva(nxt)}: {sandhi_label(row['sandhi_rule'], row['sandhi_sutra'])}")
                elif last_in_line:
                    s_cell = "पादान्ते/श्लोकान्ते अवसानम् — सन्धिः नास्ति"
                else:
                    s_cell = f"{deva(piece)} + {deva(nxt)} — सन्धिः नास्ति"
                # --- morphology + samāsa
                readings = pa.analyse(piece, scheme="iast", max_readings=3)
                best = readings[0] if readings else None
                ups, stem, stripped, conf = morph_cells(best)
                sam = row["samasa"] if len(pieces) == 1 else cs.samasa_of_piece(piece)
                if sam:
                    sam_cell = f"{' + '.join(sam['members'])} — {sam['type_guess']}"
                    if sam.get("confidence") in ("guess", "default"):
                        conf = conf + f"; samāsa-type {sam['confidence']}"
                else:
                    sam_cell = "N/A"
                alt = "; ".join(_short(r) for r in readings[1:3])
                out.append({"verse": None, "ref": None, "pada": deva(piece), "sandhi": s_cell, "upasarga": ups, "stem": stem,
                            "stripped": stripped, "samasa": sam_cell, "confidence": conf,
                            "alternatives": alt, "printed": row["word"],
                            # raw material for the Padakosha JSON (not written to xlsx/md)
                            "_reading": best, "_sam": sam, "_iast": piece})
    return out, raw


def _short(r):
    if r["type"] == "noun":
        return f"{r['stem']} {r['case_name']} {NUM_NAMES[r['num']]}"
    if r["type"] == "verb":
        return f"√{r['dhatu']} {r['lakara']} {r['purusha']}"
    return f"{r['stem']} अव्ययम्"


PERIPHRASTIC_AUX = {"āsa", "āsatuḥ", "āsuḥ", "āsitha", "āsathuḥ", "āsima", "cakāra", "cakre",
                    "cakratuḥ", "cakruḥ", "babhūva", "babhūvatuḥ", "babhūvuḥ"}


def merge_periphrastic(toks):
    """GRETIL prints the periphrastic perfect as two tokens (vivāsayām āsa); it is one word."""
    out = []
    for t in toks:
        if out and t in PERIPHRASTIC_AUX and out[-1].endswith(("ām", "āṃ")):
            out[-1] = out[-1][:-1] + "m" + t
        else:
            out.append(t)
    return out


def samhita(lines, n):
    """Continuous Devanagari: GRETIL's cosmetic space after a consonant-final word and before an
    avagraha is removed in IAST first, so conjuncts and vowel signs come out right (न्वस्मिन्)."""
    dl = []
    for line in lines:
        if not line.strip():
            continue
        x = re.sub(r"(?<=[^aāiīuūṛṝḷeoḥṃ\s'])\s+", "", line.strip())
        x = re.sub(r"\s+'", "'", x)
        dl.append(deva(x))
    return " ।\n".join(dl) + f" ॥ {n} ॥"


# ------------------------------------------------------------------ writers
def write_md(path, title, verses):
    L = [f"# {title}", "", "Generated by `sarga_analyser.py` from local tools only. The last column says how each row "
         "was obtained; «guess»/«default» rows and «pattern-only» readings should be checked by hand.", ""]
    for v in verses:
        L += [f"## श्लोकः {v['num']} ({v['ref']})", "", "**Original Verse (Devanagari):**", "",
              "> " + v["verse"].replace("\n", "  \n> "), "",
              f"**English Translation:** {v['translation']}", "",
              "| " + " | ".join(HEADERS) + " |", "|" + "---|" * len(HEADERS)]
        for r in v["rows"]:
            cells = [r["pada"], r["sandhi"], r["upasarga"], r["stem"], r["stripped"], r["samasa"], r["confidence"]]
            L.append("| " + " | ".join(str(c).replace("|", "/") for c in cells) + " |")
        L.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def write_xlsx(path, title, verses):
    import math
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sanskrit Grammatical Analysis"
    hdr_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sec_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    amber = PatternFill(start_color="FFE599", end_color="FFE599", fill_type="solid")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [22, 58, 14, 32, 58, 50, 18]
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.border = hdr_fill, Font(name="Segoe UI", size=11, bold=True, color="FFFFFF"), border
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"
    r = 2
    for v in verses:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADERS))
        c = ws.cell(row=r, column=1, value=f"श्लोकः {v['num']}  ({v['ref']})\n{v['verse']}\nEnglish Translation: {v['translation']}")
        c.font, c.fill = Font(name="Segoe UI", size=12, bold=True), sec_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 18 * (v["verse"].count("\n") + 3) + 8
        r += 1
        for row in v["rows"]:
            cells = [row["pada"], row["sandhi"], row["upasarga"], row["stem"], row["stripped"], row["samasa"], row["confidence"]]
            lines = 1
            for ci, val in enumerate(cells, 1):
                c = ws.cell(row=r, column=ci, value=val)
                c.font = Font(name="Segoe UI", size=11 if ci == 1 else 10, bold=ci == 1)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = border
                lines = max(lines, math.ceil(len(str(val)) / (widths[ci - 1] * 1.5)))
            if any(k in row["confidence"] for k in ("guess", "default", "pattern-only", "none")):
                ws.cell(row=r, column=len(HEADERS)).fill = amber
            ws.row_dimensions[r].height = 16 * lines + 6
            r += 1
        r += 1
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    wb.save(path)


def attach_verse_refs(payload, verses, kanda, sarga):
    """Add "verse_refs" (e.g. "R 1,001.003") beside the builder's "verses" list. The builder's own
    "verses" (one number per occurrence, repeats kept, keyed by the printed word) is never altered."""
    nums = sorted(v["num"] for v in verses)
    meta = payload.setdefault("corpus_meta", {})
    meta.setdefault("sarga", sarga)
    meta.setdefault("verses_covered", f"{nums[0]}-{nums[-1]}" if nums else "")
    for w in payload.get("words", []):
        w["verse_refs"] = [f"R {kanda},{sarga:03d}.{n:03d}" for n in w.get("verses", [])]
    return payload


def build_json(a, verses, verses_src, pa=None):
    """JSON payload: Padakosha layout (default) or the legacy per-verse rows."""
    if a.json_style == "rows":
        return [{**v, "rows": [{k: x for k, x in r.items() if not k.startswith("_")} for r in v["rows"]]}
                for v in verses]
    glossary = {}
    if a.glossary:
        with open(a.glossary, encoding="utf-8") as fh:
            glossary = json.load(fh)
    pb = PadakoshaBuilder(glossary, analyser=pa)
    for v in verses:                                    # verse order == order of first occurrence
        for r in v["rows"]:
            pb.add_row(v["num"], r)
    payload = pb.build(a.kanda, a.sarga, os.path.basename(a.source))
    return attach_verse_refs(payload, verses, a.kanda, a.sarga)


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", default=os.path.join(HERE, "sa_rAmAyaNa.txt"),
                    help="GRETIL-format IAST text (default: sa_rAmAyaNa.txt next to this script)")
    ap.add_argument("--kanda", type=int, required=True)
    ap.add_argument("--sarga", type=int, required=True)
    ap.add_argument("--verses", default=None, help="range, e.g. 1-10 (default: the whole sarga)")
    ap.add_argument("--out", default=None, help="output folder (default: the folder containing this script)")
    ap.add_argument("--name", default=None, help="base file name for the outputs (default: sanskrit_analysis_auto_K_SSS)")
    ap.add_argument("--formats", default="xlsx,md,json", help="any of xlsx,md,json (comma-separated)")
    ap.add_argument("--json-style", choices=("padakosha", "rows"), default="padakosha",
                    help="padakosha (default): one entry per distinct word, same layout as Ramayana_Padakosha.json; "
                         "rows: the older per-verse table dump")
    ap.add_argument("--glossary", default=None,
                    help='optional JSON file {"word or pratipadika": "English meaning", ...} used to fill meaning_en '
                         "and samasa_component_meanings in the padakosha JSON")
    ap.add_argument("--translations", default=None, help='optional JSON file {"1": "translation", ...}')
    ap.add_argument("--gold", default=None, help="folder containing a hand-written d0*.py gold analysis to score against (e.g. . for this folder)")
    a = ap.parse_args()

    for opt in ("source", "out", "translations", "gold"):
        v = getattr(a, opt)
        if v and ('"' in v or " --" in v):
            ap.error(f'--{opt} swallowed the rest of the command: {v!r}\n'
                     f'  On Windows a path ending in \\" (backslash + closing quote) escapes the quote.\n'
                     f'  Remove the trailing backslash, e.g.  --{opt} "E:\\folder\\sub"  (not "E:\\folder\\sub\\")')
    if not os.path.isfile(a.source):
        ap.error(f"--source file not found: {a.source}")
    if a.translations and not os.path.isfile(a.translations):
        ap.error(f"--translations must be a JSON file of your own translations, "
                 f'e.g. {{"1": "…", "2": "…"}} — file not found: {a.translations}  (leave the option out if you have none)')
    if a.glossary and not os.path.isfile(a.glossary):
        ap.error(f"--glossary file not found: {a.glossary}")
    if a.gold and not (os.path.isdir(a.gold) and any(f.startswith("d0") and f.endswith(".py") for f in os.listdir(a.gold))):
        ap.error(f"--gold must be a folder containing a hand-written analysis (d01.py, d02.py …) — not found in: {a.gold}"
                 "  (leave the option out if this sarga has no hand analysis)")
    out = a.out or HERE
    a.name = a.name or f"sanskrit_analysis_auto_{a.kanda}_{a.sarga:03d}"
    os.makedirs(out, exist_ok=True)
    trans = {}
    if a.translations:
        with open(a.translations, encoding="utf-8") as fh:
            trans = {int(k): v for k, v in json.load(fh).items()}
    verses_src = parse_gretil_verses(a.source, kanda=a.kanda, sarga=a.sarga)
    if not verses_src:
        sys.exit(f"No verses found for kāṇḍa {a.kanda}, sarga {a.sarga} in {a.source}")
    lo, hi = (map(int, a.verses.split("-")) if a.verses else (1, 10 ** 6))

    t0 = time.time()
    cs = CombinedSplitter()
    pa = cs._pa()
    if not pa:
        sys.exit("pada_analyser could not load its data (check that its dictionary/grammar data files are next to this script)")
    verses, raw = [], {}
    for key in sorted(verses_src):
        k, s, n = key
        if not (lo <= n <= hi):
            continue
        lines = verses_src[key]
        rows, raw[n] = analyse_verse(cs, pa, lines)
        ref = f"R {k},{s:03d}.{n:03d}"
        for r in rows:
            r["verse"], r["ref"] = n, ref
        verses.append({"num": n, "ref": ref, "verse": samhita(lines, n),
                       "translation": trans.get(n, "(not generated — no local translation engine; supply with --translations)"),
                       "rows": rows})
        print(f"  verse {n}: {len(rows)} words  ({time.time() - t0:.0f}s)", flush=True)

    title = f"Kāṇḍa {a.kanda}, Sarga {a.sarga} — automatic word-by-word analysis"
    fmts = {f.strip() for f in a.formats.split(",")}
    written = []
    if "md" in fmts:
        p = os.path.join(out, a.name + ".md"); write_md(p, title, verses); written.append(p)
    if "xlsx" in fmts:
        p = os.path.join(out, a.name + ".xlsx"); write_xlsx(p, title, verses); written.append(p)
    if "json" in fmts:
        p = os.path.join(out, a.name + ".json")
        with open(p, "w", encoding="utf-8") as fh:
            json.dump(build_json(a, verses, verses_src, pa), fh, ensure_ascii=False, indent=1)
        written.append(p)
    n_rows = sum(len(v["rows"]) for v in verses)
    flagged = sum(1 for v in verses for r in v["rows"] if any(x in r["confidence"] for x in ("guess", "default", "pattern-only", "none")))
    print(f"\n{len(verses)} verses, {n_rows} words ({flagged} rows flagged for checking) in {time.time() - t0:.0f}s")
    for p in written:
        print("  wrote", p)

    if a.gold:
        sysp = os.path.join(out, a.name + "_splitter_raw.json")
        with open(sysp, "w", encoding="utf-8") as fh:
            json.dump(raw, fh, ensure_ascii=False)
        import subprocess
        morph = os.path.join(out, a.name + "_morph.json")
        subprocess.run([sys.executable, os.path.join(HERE, "run_morphology.py"), sysp, "--out", morph], check=True)
        subprocess.run([sys.executable, os.path.join(HERE, "compare_with_gold.py"), "--system", sysp, "--morph", morph,
                        "--gold", a.gold, "--out-md", os.path.join(out, a.name + "_scorecard.md"),
                        "--out-json", os.path.join(out, a.name + "_comparison.json")], check=True)
        print("  wrote", os.path.join(out, a.name + "_scorecard.md"))


if __name__ == "__main__":
    main()
