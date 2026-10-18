#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_sarga1_report_v2_xlsx.py
===================================
Excel version of generate_sarga1_report_v2.py's data. Sandhi columns are
considered verified (fixed, guard-tested, spot-checked against the exact
words a human reviewer flagged wrong in the original run). Samasa
columns are substantially hardened since the first pass of this session
(258 -> 189 decompositions, after two new guards: rejecting
samasa_splitter.py's own self-flagged "NOT a samasa" particle-coincidence
candidates, and requiring every member to independently pass a STRICTER
dictionary check than samasa_splitter's own permissive one) but are
still NOT fully verified -- a follow-up full audit of the 189 remaining
decompositions found most nominal compounds now correct, but two
residual failure categories:
  1. VERB FORMS (perfect tense, passive participles, gerunds/causatives
     not covered by samasa_splitter's looks_like_verb_form table) still
     get wrongly torn into root+ending fragments, e.g. ददर्श (perfect of
     दृश्) -> दद+ऋश, निजघान (perfect of हन्) -> निज्+अघ+अन.
  2. A short, genuinely real but RARE dictionary word (इन "the sun",
     अणि, तस्, ...) coincidentally matches a stripped case-ending
     fragment, e.g. धनदेन (single word, instrumental of धनद) -> wrongly
     धनद + इन.
Both need broader verb-paradigm / frequency-aware work beyond this
session's guards -- flagged in the Notes sheet, not hidden.

USAGE
-----
    D:\\python\\python.exe generate_sarga1_report_v2_xlsx.py [output.xlsx]
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SANDHI_DIR = os.path.normpath(os.path.join(HERE, "..", "sandhi"))
for d in (SANDHI_DIR, HERE):
    if d not in sys.path:
        sys.path.insert(0, d)

import openpyxl  # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from generate_sarga1_report_v2 import build_report_rows  # noqa: E402
from combined_splitter import CombinedSplitter  # noqa: E402

DEFAULT_OUTPUT = os.path.join(HERE, "sarga1_sandhi_samasa_report_v2.xlsx")

HEADERS = ["#", "Word", "Verse", "Sandhi Rule(s)", "Sandhi Pada-Vigraha",
           "Samasa Type (IMPROVED, still verify)", "Samasa Pada-Vigraha (IMPROVED, still verify)",
           "Verdict (mark here)", "Correction / Notes"]
COL_WIDTHS = [5, 22, 14, 20, 30, 30, 30, 16, 40]

SAMASA_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def main():
    output_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUTPUT

    cs = CombinedSplitter()
    rows, n_sandhi, n_samasa = build_report_rows(cs)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sarga 1 report (v2)"

    ws.append(HEADERS)
    for col, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for i, row in enumerate(rows, 1):
        ws.append([i, row["word"], row["verse"], row["sandhi_name"], row["sandhi_vigraha"],
                   row["samasa_name"], row["samasa_vigraha"], "", ""])
        # Highlight the two samasa columns (F, G) light amber, but only
        # where a decomposition actually fired -- marks "still verify
        # this" without visually flooding all 880 rows (most of which
        # are blank "-" in these columns anyway).
        if row["samasa_name"] != "-":
            ws.cell(row=i + 1, column=6).fill = SAMASA_FILL
            ws.cell(row=i + 1, column=7).fill = SAMASA_FILL

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column > 3))

    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Valmiki Ramayana -- Balakanda, Sarga 1 (R_1,001.001-079) -- v2"
    ws2["A2"] = "Combined sandhi_splitter.py + samasa_splitter.py analysis"
    ws2["A3"] = (f"Total words: {len(rows)}  |  Sandhi restorations found: {n_sandhi}  |  "
                 f"Samasa decompositions found: {n_samasa}")
    ws2["A5"] = (
        "SANDHI COLUMNS: fixed and verified this session. v1's report checked sandhi by "
        "running sandhi_splitter.py's internal-word boundary finder on each ISOLATED GRETIL "
        "token one at a time -- wrong, because it manufactures a fake internal cut even in a "
        "single already-complete word with no real junction (confirmed by direct human review: "
        "~90%+ of v1's Sandhi Rule(s) column was wrong). v2 instead checks sandhi PAIRWISE, "
        "between each two adjacent words exactly as the source text (GRETIL) already segments "
        "them, via samasa_splitter.py's restore_word_boundary (the real forward Sandhi.sandhi() "
        "engine as oracle, cross-checked against apte_dictionary.sqlite). Spot-checked directly "
        "against the specific words flagged wrong in v1 (परिपप्रच्छ, वाल्मीकिर्, नारदं, धर्मज्ञश्, "
        "सत्यवाक्यो, दृढव्रतः, चारित्रेण, सर्वभूतेषु, ...) and all now resolve correctly."
    )
    ws2["A7"] = (
        "SAMASA COLUMNS: substantially hardened this session (258 -> 189 decompositions across "
        "the same 880 words) but still NOT fully verified -- highlighted light amber only on "
        "rows where a decomposition fired. Two guards added: (1) reject any candidate "
        "samasa_splitter.py's OWN guess_samasa_type() already labels 'NOT a samasa' (a particle "
        "coincidentally matching the start of an ordinary word, e.g. चारित्रेण was wrongly shown "
        "as च+अरित्रेण); (2) require every member to independently pass a STRICTER dictionary "
        "check (sandhi_splitter.py's exact/H/M-strip match, or its small hand-curated "
        "common-inflection table) instead of trusting samasa_splitter's own looser 26-suffix, "
        "two-rounds-deep declension guesser, which had been accepting coincidental matches "
        "(fixed: विष्णुना, धैर्येण, नियतात्मा, अदीनात्मा, इवापरः, प्रीत्या, and more)."
    )
    ws2["A9"] = (
        "STILL OPEN, two categories found by a follow-up full audit of all 189 remaining rows: "
        "(1) VERB FORMS -- perfect tense, passive participles, and gerund/causative patterns "
        "samasa_splitter.py's looks_like_verb_form table doesn't cover yet get wrongly torn into "
        "root+ending fragments (ददर्श -> दद+ऋश, निजघान -> निज्+अघ+अन, वार्यमाणः -> वारि+अम+अणः). "
        "(2) a short but genuinely real, RARE dictionary word (इन 'the sun', अणि, तस्, ...) "
        "coincidentally matches a stripped case-ending fragment even under the stricter check "
        "(धनदेन, single word instrumental of धनद, wrongly shown as धनद+इन). Both need broader "
        "verb-paradigm coverage / a frequency-aware check, not another one-off guard -- flagging "
        "here rather than continuing to whack-a-mole individual cases."
    )
    for cell in ("A5", "A7", "A9"):
        ws2[cell].alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 110
    ws2["A11"] = "Reviewer columns: mark 'Verdict' (e.g. OK / WRONG / PARTIAL) and add notes in 'Correction / Notes' for BOTH sandhi and samasa now."

    wb.save(output_path)
    print(f"\nWrote {output_path} ({len(rows)} words, "
          f"{n_sandhi} sandhi restorations, {n_samasa} samasa decompositions [not yet reliable])",
          file=sys.stderr)


if __name__ == "__main__":
    main()
